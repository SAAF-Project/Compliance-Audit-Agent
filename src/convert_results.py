#%%
import json
import glob
import os

from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT

from validate_func import *

SEL_FILES = '20251215_165752'


def get_heading_style(doc: Document, preferred: str = "Heading 1"):
    """
    Try to get a heading style robustly across localized Word templates.
    Falls back to any style containing 'Heading' (case-insensitive),
    otherwise returns 'Normal'.
    """
    # Try preferred first
    try:
        return doc.styles[preferred]
    except KeyError:
        pass

    # Try common localized names
    for candidate in ["Heading 1", "Kop 1", "Titre 1", "Encabezado 1"]:
        try:
            return doc.styles[candidate]
        except KeyError:
            continue

    # Fallback: first style that contains 'Heading'
    for s in doc.styles:
        if "heading" in s.name.lower():
            return s

    # Last resort
    return doc.styles["Normal"]

def is_list_of_dicts_with_uniform_keys(value):
    """
    Returns (True, keys) if value is a list of dicts with uniform keys, else (False, None).
    If list is empty, treat as uniform with empty keys (will render an empty table with header only).
    """
    if not isinstance(value, list):
        return False, None
    dicts = [item for item in value if isinstance(item, dict)]
    if len(dicts) != len(value):
        return False, None
    if len(dicts) == 0:
        return True, []
    # Check uniform keys across items (use set equality)
    keyset = set(dicts[0].keys())
    for d in dicts[1:]:
        if set(d.keys()) != keyset:
            return False, None
    return True, list(dicts[0].keys())

def add_dict_as_one_row_table(doc: Document, d: dict, include_header_row: bool = True):
    """
    Render a dict as a table. If include_header_row=True, create a header + one data row.
    If False, create a single data row only.
    """
    if not isinstance(d, dict):
        # Safety: write as text if not a dict
        p = doc.add_paragraph(str(d))
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return

    keys = list(d.keys())
    if include_header_row:
        rows = 2
        table = doc.add_table(rows=rows, cols=len(keys))
        table.style = "Table Grid"
        # Header
        hdr = table.rows[0].cells
        for idx, k in enumerate(keys):
            hdr[idx].text = str(k)
        # Data
        row = table.rows[1].cells
        for idx, k in enumerate(keys):
            if idx == 0:
               row[idx].text = "" if d.get(k) is None else str(d.get(k)).split(':')[0]
            else:
               row[idx].text = "" if d.get(k) is None else str(d.get(k))
    else:
        rows = 1
        table = doc.add_table(rows=rows, cols=len(keys))
        table.style = "Table Grid"
        row = table.rows[0].cells
        for idx, k in enumerate(keys):
            row[idx].text = "" if d.get(k) is None else str(d.get(k))

def add_list_of_dicts_as_table(doc: Document, items: list, keys: list):
    """
    Render a list of dicts (uniform keys) as a multi-row table with a header.
    Handles missing values by blank cells.
    """
    n_rows = 1 + (len(items))  # header + one row per item
    table = doc.add_table(rows=n_rows, cols=len(keys))
    table.style = "Table Grid"

    # Header
    hdr = table.rows[0].cells
    for col, k in enumerate(keys):
        hdr[col].text = str(k)

    # Data rows
    for r, item in enumerate(items, start=1):
        row_cells = table.rows[r].cells
        for c, k in enumerate(keys):
            row_cells[c].text = "" if item.get(k) is None else str(item.get(k))

def add_scalar_or_list_value(doc: Document, value):
    """
    Render scalars as text; list of scalars as bullets; anything else as JSON string.
    """
    if isinstance(value, (str, int, float, bool)) or value is None:
        doc.add_paragraph("" if value is None else str(value))
    elif isinstance(value, list) and all(not isinstance(x, dict) for x in value):
        for x in value:
            run = doc.add_paragraph(style=None).add_run("" if x is None else str(x))
            run.text = "- " + run.text
    else:
        doc.add_paragraph(json.dumps(value, ensure_ascii=False, indent=2))


def _set_cell_lines(cell, lines):
    """Write a list of strings as separate paragraphs inside a table cell."""
    str_lines = [str(l) for l in lines if l is not None]
    if not str_lines:
        cell.text = ""
        return
    cell.paragraphs[0].text = str_lines[0]
    for line in str_lines[1:]:
        cell.add_paragraph(line)


def _format_timestamp(ts) -> str:
    """Return only the date part of an ISO timestamp string."""
    return str(ts).split("T")[0] if ts else ""


def _prettify_gap_key(key: str) -> str:
    """Convert 'a_ability_to_perform_and_manage' -> '(a) Ability to perform and manage'."""
    parts = key.split("_", 1)
    if len(parts) == 2 and len(parts[0]) == 1 and parts[0].isalpha():
        return f"({parts[0]}) {parts[1].replace('_', ' ').capitalize()}"
    return key.replace("_", " ").capitalize()


def _stringify_ref_item(item) -> str:
    """Flatten a single Policy Source Reference (string or dict) into one line."""
    if isinstance(item, dict):
        source = str(item.get("Source", "")).strip()
        sections = item.get("Key Sections", "")
        if isinstance(sections, list):
            sections = ", ".join(str(s) for s in sections if s)
        ts = _format_timestamp(item.get("Timestamp", ""))
        return " - ".join(p for p in (source, str(sections).strip(), ts) if p)
    return "" if item is None else str(item)


def add_policy_source_refs_list(doc: Document, value):
    """Render Policy Source References as a bullet list (not a table). Accepts a
    plain string, a list of strings, or a legacy list of dicts."""
    if isinstance(value, list):
        add_scalar_or_list_value(doc, [_stringify_ref_item(it) for it in value])
    else:
        add_scalar_or_list_value(doc, value)


def add_gap_output_section(doc: Document, data: dict, heading_style):
    """Render gap_output (canonical schema) as sub-tables, mirroring how the
    top-level compliance keys are rendered. Each sub-key (Generalized Gap Audit,
    Detailed Gap Audit, Policy Source References, Audit Summary) gets a bold
    sub-heading and an appropriate table / list / paragraph."""
    if not isinstance(data, dict):
        add_scalar_or_list_value(doc, data)
        return

    skip = {"Guideline Items", "Guideline Title"}
    for key, value in data.items():
        if key in skip:
            continue
        p = doc.add_paragraph(str(key))
        p.runs[0].bold = True

        if key == "Policy Source References":
            add_policy_source_refs_list(doc, value)
        elif isinstance(value, dict):
            add_dict_as_one_row_table(doc, value, include_header_row=True)
        elif isinstance(value, list):
            is_uniform, keys = is_list_of_dicts_with_uniform_keys(value)
            if is_uniform and keys:
                add_list_of_dicts_as_table(doc, value, keys)
            else:
                add_scalar_or_list_value(doc, value)
        else:
            add_scalar_or_list_value(doc, value)


def add_reporter_output_section(doc: Document, data: dict):
    """Render the structured reporter_output: summary paragraph + one table row per recommendation."""
    # Summary intro
    summary = data.get("summary", "")
    if summary:
        doc.add_paragraph(str(summary))

    recommendations = data.get("recommendations", [])
    if not recommendations:
        # Fallback for legacy/unstructured output
        if not summary:
            add_scalar_or_list_value(doc, data)
        return

    cols = ["#", "Title", "Guideline Reference", "Gap",
            "Target Policy", "Insertion Point", "Recommended Wording", "Rationale"]
    table = doc.add_table(rows=1 + len(recommendations), cols=len(cols))
    table.style = "Table Grid"
    table.autofit = True
    table.allow_autofit = True
    for c, k in enumerate(cols):
        table.rows[0].cells[c].text = k

    for r, rec in enumerate(recommendations, start=1):
        cells = table.rows[r].cells
        cells[0].text = str(rec.get("id", r))
        cells[1].text = str(rec.get("title", ""))
        cells[2].text = str(rec.get("guidelineReference", ""))
        cells[3].text = str(rec.get("gap", ""))
        cells[4].text = str(rec.get("targetPolicy", ""))
        cells[5].text = str(rec.get("insertionPoint", ""))
        wording = rec.get("recommendedWording", "")
        _set_cell_lines(cells[6], str(wording).split("\n"))
        cells[7].text = str(rec.get("rationale", ""))


def json_to_word(
    input_json_path: str,
    output_docx_path: str = "output.docx",
    heading_style_name: str = "Heading 1",
    include_header_for_one_row_dict: bool = True,
    model_config_description: str = '',
    val_summary: list = [],
):
    """
    Convert a JSON into a Word document:
    - Each top-level key becomes a Heading 1 section (chapter head).
    - Dict values -> one-row table (optionally with header).
    - List of dicts (uniform keys) -> table with header + multiple rows.
    - Others -> text or bullet list.
    """
    input_path = Path(input_json_path)
    if not input_path.exists():
        raise FileNotFoundError(f"JSON not found: {input_path}")

    with input_path.open("r", encoding="utf-8") as f:
        data = json.load(f)

    if not isinstance(data, dict):
        raise ValueError("Top-level JSON must be an object (dict).")

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    heading_style = get_heading_style(doc, preferred=heading_style_name)

    # Optionally set a base font for Normal style for readability
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # Add top header and sentence before the loop
    doc.add_paragraph("Model config", style=heading_style)  # Header 1
    doc.add_paragraph(model_config_description)

    _HEADING_ALIASES = {
        "selector_output": "Policy Selections",
        "gap_output": "Compliance Gap Testing",
        "reporter_output": "Reporter Output",
    }

    # Preserve the insertion order of JSON keys
    for top_key, value in data.items():
        display_key = _HEADING_ALIASES.get(top_key, top_key)
        p = doc.add_paragraph(str(display_key))
        p.style = heading_style

        if top_key == "selector_output" and isinstance(value, dict):
            policies = value.get("policies", [])
            is_uniform, keys = is_list_of_dicts_with_uniform_keys(policies)
            if is_uniform and keys:
                add_list_of_dicts_as_table(doc, policies, keys)
            else:
                add_scalar_or_list_value(doc, policies)
        elif top_key == "gap_output" and isinstance(value, dict):
            add_gap_output_section(doc, value, heading_style)
        elif top_key == "reporter_output":
            if isinstance(value, dict):
                add_reporter_output_section(doc, value)
            else:
                doc.add_paragraph(str(value) if value else "")
        elif top_key == "Policy Source References":
            add_policy_source_refs_list(doc, value)
        elif isinstance(value, dict):
            add_dict_as_one_row_table(doc, value, include_header_row=include_header_for_one_row_dict)
        else:
            is_uniform, keys = is_list_of_dicts_with_uniform_keys(value)
            if is_uniform:
                add_list_of_dicts_as_table(doc, value, keys)
            else:
                add_scalar_or_list_value(doc, value)

    if len(val_summary) > 0:
       doc.add_paragraph("Validation Summary", style=heading_style)  # Header 1
       add_list_of_dicts_as_table(doc, val_summary, list(val_summary[0].keys()))
       doc.add_paragraph("")

    doc.save(output_docx_path)
    return output_docx_path

def _with_line_numbers(lines: list, start: int = 0) -> str:
    width = max(2, len(str(start + len(lines) - 1)))
    return "\n".join(f"{str(i).rjust(width)} | {line}" for i, line in enumerate(lines, start=start))


if __name__ == "__main__":
    all_policies = glob.glob(str(INPUT_FOLDER / "plain_docs" / "*.txt"))
    if LIMIT_FILES > 0:
        all_policies = all_policies[:LIMIT_FILES]

    mdpolicies = []
    for i, policy_path in enumerate(all_policies):
        with open(policy_path, "r", encoding="utf-8", errors="ignore") as f:
            mdpolicy = f.read().splitlines()
            numbered_policies = _with_line_numbers(mdpolicy)

            letter = chr(65 + i)
            mdpolicies.extend([f"--- Policy {letter}: {os.path.basename(policy_path)} ---\n"])
            mdpolicies.extend(numbered_policies.split('\n'))

    if SEL_FILES != 'all':
       if type(SEL_FILES) is not list:
         SEL_FILES = [SEL_FILES]

       for timestamp in SEL_FILES:
         json_path = f'./prog_res/output_{timestamp}.json'

         with open(json_path, 'r', encoding='utf-8') as f:
            data = json.load(f)

         val_summary = write_val_res(data.get('Detailed Compliance Audit', {}), mdpolicies, timestamp)

         out = json_to_word(
               input_json_path=json_path,
               output_docx_path=f"./prog_res/compliance_audit_gdl_{timestamp}.docx",
               heading_style_name="Heading 1",
               include_header_for_one_row_dict=True,
               val_summary=val_summary,
         )
         print(f"Processed and saved to {out}")

    else:
       all_files = glob.glob('./prog_res/output_*.json')
       for json_path in all_files:
         if 'raw' in json_path:
            continue
         timestamp = Path(json_path).stem.split('output_')[-1]
         out = json_to_word(
            input_json_path=json_path,
            output_docx_path=f"./prog_res/compliance_audit_gdl_{timestamp}.docx",
            heading_style_name="Heading 1",
            include_header_for_one_row_dict=True,
         )
         print(f"Processed and saved to {out}")
